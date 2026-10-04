"""Extraction from attached documents.

The property under test throughout is the one in the module docstring: an
attachment that was *not* read must never be reportable as clean. Every case
below is either "we read it, here is the text" or "we did not read it, here is
why", and `Extraction.inspected` is what separates them.
"""

from __future__ import annotations

import io
import shutil
import zipfile

import pytest
from legacy_office import FIXTURE, FIXTURE_TEXT, build_cfb, doc_bytes, ppt_bytes, xls_bytes
from llmp_redaction.documents import MAX_BYTES, extract, media_type_for

# markitdown is services/redaction's dependency, not the gateway's, and the root
# suite deliberately runs this directory (pyproject's `testpaths`) so these tests
# do not rot in a corner nobody executes. In the gateway's environment the
# extractor's library is absent, so these skip — visibly, with a reason — and
# run for real under `uv run --python 3.13 pytest` in services/redaction, which
# is where the 37-package tree is installed.
pytest.importorskip(
    "markitdown",
    reason="document extraction needs services/redaction's own environment",
)
from llmp_shared import ExtractionKind, ExtractionResponse

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"

IBAN = "IT60X0542811101000000123456"


def docx_bytes(paragraph: str = "", table: list[list[str]] | None = None) -> bytes:
    import docx

    document = docx.Document()
    if paragraph:
        document.add_paragraph(paragraph)
    if table:
        added = document.add_table(rows=len(table), cols=len(table[0]))
        for row_index, row in enumerate(table):
            for cell_index, value in enumerate(row):
                added.cell(row_index, cell_index).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def xlsx_bytes(rows: list[list[object]]) -> bytes:
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def pptx_bytes(title: str, notes: str = "") -> bytes:
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = title
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def pdf_bytes(text: str | None) -> bytes:
    """A minimal one-page PDF. `text=None` produces a page with no text layer,
    which is what a scan is once the image is stripped out."""
    if text is None:
        content = b"q Q"  # a valid, empty content stream: a blank page
    else:
        escaped = text.replace("(", r"\(").replace(")", r"\)").encode()
        content = b"BT /F1 12 Tf 72 720 Td (" + escaped + b") Tj ET"

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += str(number).encode() + b" 0 obj\n" + body + b"\nendobj\n"
    start = len(out)
    out += b"xref\n0 " + str(len(objects) + 1).encode() + b"\n0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        b"trailer\n<< /Size "
        + str(len(objects) + 1).encode()
        + b" /Root 1 0 R >>\nstartxref\n"
        + str(start).encode()
        + b"\n%%EOF\n"
    )
    return bytes(out)


class TestReadsWhatIsThere:
    def test_a_word_document(self) -> None:
        result = extract(docx_bytes(f"Please invoice {IBAN}."), media_type=DOCX)
        assert result.kind is ExtractionKind.TEXT
        assert IBAN in result.text
        assert result.extractor == "markitdown"

    def test_a_word_table(self) -> None:
        """An invoice *is* a table, and tables are not in `document.paragraphs` —
        reading only paragraphs would miss where the personal data actually is."""
        data = docx_bytes(table=[["Customer", "Luca Bianchi"], ["IBAN", IBAN]])
        result = extract(data, media_type=DOCX)
        assert result.inspected
        assert "Luca Bianchi" in result.text
        assert IBAN in result.text

    def test_a_spreadsheet(self) -> None:
        data = xlsx_bytes([["name", "iban"], ["Luca Bianchi", IBAN]])
        result = extract(data, media_type=XLSX)
        assert result.inspected
        assert IBAN in result.text
        assert "Luca Bianchi" in result.text

    def test_presentation_notes_are_read(self) -> None:
        """Speaker notes are text someone wrote and forgot about, which makes
        them a likelier home for personal data than the slide."""
        result = extract(pptx_bytes("Quarterly review", notes=f"chase {IBAN}"), media_type=PPTX)
        assert result.inspected
        assert IBAN in result.text

    def test_a_text_layer_pdf(self) -> None:
        result = extract(pdf_bytes(f"Invoice for {IBAN}"), media_type="application/pdf")
        assert result.kind is ExtractionKind.TEXT
        assert IBAN in result.text
        assert result.extractor == "markitdown"


class TestNotReadingIsNotCleanliness:
    """The property this module exists for."""

    def test_a_pdf_with_no_text_layer_is_not_reported_as_clean(self) -> None:
        """A scan is a picture of a page. Extraction finds nothing, and nothing
        must not be indistinguishable from a document containing nothing."""
        result = extract(pdf_bytes(None), media_type="application/pdf")
        assert result.kind is ExtractionKind.NO_TEXT_LAYER
        assert result.text == ""
        assert not result.inspected

    def test_an_unsupported_type_says_so(self) -> None:
        result = extract(b"\x00\x01\x02", media_type="image/png", filename="scan.png")
        assert result.kind is ExtractionKind.UNSUPPORTED
        assert not result.inspected

    def test_a_corrupt_document_is_unreadable_not_empty(self) -> None:
        result = extract(b"PK\x03\x04 not really a docx", media_type=DOCX)
        assert result.kind is ExtractionKind.UNREADABLE
        assert not result.inspected

    def test_an_oversized_document_is_refused_before_parsing(self) -> None:
        result = extract(b"x" * (MAX_BYTES + 1), media_type="application/pdf")
        assert result.kind is ExtractionKind.TOO_LARGE
        assert not result.inspected

    def test_a_zip_bomb_is_refused_on_its_declared_size(self) -> None:
        """Read the central directory, not the payload: the whole point of a zip
        bomb is that it is small until you decompress it."""
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("word/document.xml", b"\x00" * (200 * 1024 * 1024))
        result = extract(buffer.getvalue(), media_type=DOCX)
        assert result.kind is ExtractionKind.UNREADABLE
        assert "declares" in result.detail

    @pytest.mark.parametrize("kind", list(ExtractionKind))
    def test_only_text_counts_as_inspected(self, kind: ExtractionKind) -> None:
        assert ExtractionResponse(kind=kind).inspected is (kind is ExtractionKind.TEXT)


class TestTypeResolution:
    def test_a_declared_type_is_used(self) -> None:
        assert media_type_for(DOCX, None) == DOCX

    def test_an_octet_stream_falls_back_to_the_extension(self) -> None:
        """Several clients send `application/octet-stream` for any upload."""
        assert media_type_for("application/octet-stream", "invoice.docx") == DOCX
        assert media_type_for(None, "report.pdf") == "application/pdf"

    def test_a_name_never_overrides_a_type_we_know(self) -> None:
        """A PDF named `.docx` is parser confusion waiting to happen."""
        assert media_type_for("application/pdf", "invoice.docx") == "application/pdf"

    def test_an_unknown_type_and_name_resolve_to_nothing(self) -> None:
        assert media_type_for("application/x-thing", "data.bin") is None


TEXTBOX_RUN = (
    '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
    'xmlns:v="urn:schemas-microsoft-com:vml" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing">'
    "<w:r><mc:AlternateContent>"
    '<mc:Choice Requires="wps"><w:drawing><wp:anchor><wps:wsp><wps:txbx><w:txbxContent>'
    "<w:p><w:r><w:t>{text}</w:t></w:r></w:p>"
    "</w:txbxContent></wps:txbx></wps:wsp></wp:anchor></w:drawing></mc:Choice>"
    "<mc:Fallback><w:pict><v:shape><v:textbox><w:txbxContent>"
    "<w:p><w:r><w:t>{text}</w:t></w:r></w:p>"
    "</w:txbxContent></v:textbox></v:shape></w:pict></mc:Fallback>"
    "</mc:AlternateContent></w:r></w:p>"
)


def docx_with_textbox(body: str, boxed: str) -> bytes:
    """A real .docx whose second paragraph holds a text box, written the way
    Word 2010+ writes one: a DrawingML choice with a VML fallback."""
    source = zipfile.ZipFile(io.BytesIO(docx_bytes(body)))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == "word/document.xml":
                xml = content.decode("utf-8")
                xml = xml.replace("<w:sectPr", TEXTBOX_RUN.format(text=boxed) + "<w:sectPr", 1)
                content = xml.encode("utf-8")
            target.writestr(item, content)
    return out.getvalue()


class TestOOXMLStillRead:
    def test_a_text_box_does_not_stop_the_body_being_read(self) -> None:
        data = docx_with_textbox("Corpo del documento", "Nel riquadro")
        result = extract(data, media_type=DOCX)
        assert result.kind is ExtractionKind.TEXT
        assert "Corpo del documento" in result.text

    def test_a_zip_is_read_by_its_parts_not_its_label(self) -> None:
        """A real .docx sent as `application/msword` is still a .docx."""
        result = extract(docx_bytes(f"Invoice {IBAN}"), media_type="application/msword")
        assert result.kind is ExtractionKind.TEXT
        assert IBAN in result.text

    def test_a_corrupt_docx_names_what_was_wrong(self) -> None:
        result = extract(b"definitely not office", media_type=DOCX)
        assert result.kind is ExtractionKind.UNREADABLE
        assert ".docx" in result.detail and "not an Office document" in result.detail


class TestLegacyOffice:
    """Pre-2007 files are OLE2 compound files, not zips. The container is read
    from the bytes, so a label cannot send one to the wrong reader."""

    def test_the_committed_fixture_is_what_the_generator_writes(self) -> None:
        assert FIXTURE.read_bytes() == doc_bytes(FIXTURE_TEXT)

    def test_a_legacy_word_document_is_read_by_antiword(self) -> None:
        if shutil.which("antiword") is None:
            pytest.skip("antiword is installed in the extractor image, not on this host")
        result = extract(FIXTURE.read_bytes(), media_type="application/msword")
        assert result.kind is ExtractionKind.TEXT
        assert result.extractor == "antiword"
        assert IBAN in result.text
        assert "Luca Bianchi" in result.text

    def test_italian_accents_survive(self) -> None:
        if shutil.which("antiword") is None:
            pytest.skip("antiword is installed in the extractor image, not on this host")
        result = extract(FIXTURE.read_bytes(), media_type="application/msword")
        assert "àèéìòù ÀÈÉÌÒÙ" in result.text
        assert "Menù" in result.text and "Perché è già così" in result.text

    def test_the_same_bytes_declared_docx_are_still_a_legacy_doc(self) -> None:
        """The reported case: `MENU ... .doc.docx`."""
        if shutil.which("antiword") is None:
            pytest.skip("antiword is installed in the extractor image, not on this host")
        result = extract(
            FIXTURE.read_bytes(),
            media_type=DOCX,
            filename="MENU' autunno inverno non vidimato 2021 2022.doc.docx",
        )
        assert result.kind is ExtractionKind.TEXT
        assert result.extractor == "antiword"
        assert "caffè" in result.text

    def test_the_label_does_not_matter_even_when_missing(self) -> None:
        if shutil.which("antiword") is None:
            pytest.skip("antiword is installed in the extractor image, not on this host")
        result = extract(FIXTURE.read_bytes(), media_type="application/octet-stream")
        assert result.extractor == "antiword"

    def test_lines_are_not_wrapped(self) -> None:
        if shutil.which("antiword") is None:
            pytest.skip("antiword is installed in the extractor image, not on this host")
        long_line = "parola " * 40 + IBAN
        data = doc_bytes(long_line)
        result = extract(data, media_type="application/msword")
        assert IBAN in result.text and "\n" not in result.text.strip()

    def test_without_antiword_the_reason_says_so(self, monkeypatch: pytest.MonkeyPatch) -> None:
        monkeypatch.setattr(shutil, "which", lambda _name: None)
        result = extract(FIXTURE.read_bytes(), media_type="application/msword")
        assert result.kind is ExtractionKind.UNSUPPORTED
        assert not result.inspected
        assert result.detail == "Old Word .doc files need antiword in the extractor image"

    def test_antiword_that_hangs_is_cut_off(self, monkeypatch: pytest.MonkeyPatch) -> None:
        import subprocess

        def hang(*_args: object, **_kwargs: object) -> None:
            raise subprocess.TimeoutExpired("antiword", 1)

        monkeypatch.setattr(shutil, "which", lambda _name: "/usr/bin/antiword")
        monkeypatch.setattr(subprocess, "run", hang)
        result = extract(FIXTURE.read_bytes(), media_type="application/msword")
        assert result.kind is ExtractionKind.UNREADABLE
        assert "took longer" in result.detail

    def test_a_legacy_excel_file_is_still_read(self) -> None:
        data = xls_bytes([["nome", "iban"], ["Luca Bianchi", f"{IBAN} è"]])
        result = extract(data, media_type="application/vnd.ms-excel")
        assert result.kind is ExtractionKind.TEXT
        assert IBAN in result.text and "Luca Bianchi" in result.text

    def test_a_legacy_excel_file_labelled_xlsx_is_still_read(self) -> None:
        result = extract(xls_bytes([["a", "Luca Bianchi"]]), media_type=XLSX, filename="x.xlsx")
        assert result.kind is ExtractionKind.TEXT
        assert "Luca Bianchi" in result.text

    def test_a_legacy_powerpoint_file_is_refused_by_name(self) -> None:
        result = extract(ppt_bytes(), media_type="application/vnd.ms-powerpoint")
        assert result.kind is ExtractionKind.UNSUPPORTED
        assert not result.inspected
        assert ".ppt" in result.detail and "save as .pptx" in result.detail

    def test_a_legacy_powerpoint_file_labelled_pptx_is_refused_by_name(self) -> None:
        result = extract(ppt_bytes(), media_type=PPTX)
        assert ".ppt" in result.detail

    def test_a_password_protected_office_file_says_so(self) -> None:
        locked = build_cfb({"EncryptedPackage": b"x", "EncryptionInfo": b"x"})
        result = extract(locked, media_type=DOCX)
        assert result.kind is ExtractionKind.UNSUPPORTED
        assert "password-protected" in result.detail

    def test_an_unknown_compound_file_is_refused(self) -> None:
        result = extract(build_cfb({"Something": b"x"}), media_type=DOCX)
        assert result.kind is ExtractionKind.UNSUPPORTED
        assert "OLE" in result.detail

    def test_a_truncated_compound_file_does_not_raise(self) -> None:
        result = extract(FIXTURE.read_bytes()[:700], media_type=DOCX)
        assert not result.inspected
        assert result.detail


def scan_pdf(pages: int, size: tuple[int, int] = (2000, 3000)) -> bytes:
    """An image-only PDF: what a scanner produces, with no text layer."""
    from PIL import Image, ImageDraw

    images = []
    for number in range(pages):
        image = Image.new("RGB", size, "white")
        ImageDraw.Draw(image).rectangle((100, 100, size[0] - 100, 400 + number), fill="black")
        images.append(image)
    buffer = io.BytesIO()
    images[0].save(buffer, format="PDF", save_all=True, append_images=images[1:])
    return buffer.getvalue()


LONG_TEXT = "Il presente documento contiene del testo vero, non una fotografia di testo."


def mixed_pdf(*parts: bytes) -> bytes:
    """PDFs joined in order, so a typed page can sit among scanned ones."""
    import pypdfium2 as pdfium

    merged = pdfium.PdfDocument.new()
    for part in parts:
        merged.import_pages(pdfium.PdfDocument(part))
    buffer = io.BytesIO()
    merged.save(buffer)
    return buffer.getvalue()


class TestPdfPagesAsImages:
    def test_a_scan_without_the_request_is_unchanged(self) -> None:
        result = extract(scan_pdf(2), media_type="application/pdf")
        assert result.kind is ExtractionKind.NO_TEXT_LAYER
        assert result.page_images == []

    def test_a_scan_returns_jpegs_scaled_to_the_long_side(self) -> None:
        import base64

        from PIL import Image

        result = extract(scan_pdf(2), media_type="application/pdf", page_images=20)
        # Still "not read": the pictures are for a model, not text.
        assert result.kind is ExtractionKind.NO_TEXT_LAYER
        assert not result.inspected
        assert [image.page for image in result.page_images] == [1, 2]
        assert result.pages == 2 and result.images_truncated is False
        assert result.page_texts == ["", ""]
        picture = Image.open(io.BytesIO(base64.b64decode(result.page_images[0].data)))
        assert picture.format == "JPEG"
        assert max(picture.size) == 1280  # the default long side
        assert result.page_images[0].mime == "image/jpeg"

    def test_the_long_side_is_the_callers_and_is_clamped(self) -> None:
        import base64

        from PIL import Image

        def longest(requested: int) -> int:
            result = extract(
                scan_pdf(1), media_type="application/pdf", page_images=1, long_side=requested
            )
            data = base64.b64decode(result.page_images[0].data)
            return int(max(Image.open(io.BytesIO(data)).size))

        assert longest(900) == 900
        assert longest(99999) == 2048

    def test_each_page_is_judged_on_its_own(self) -> None:
        """A typed cover sheet and two scanned pages: text for the first, pictures
        for the others, in page order."""
        pdf = mixed_pdf(pdf_bytes(LONG_TEXT), scan_pdf(2, (200, 300)))
        result = extract(pdf, media_type="application/pdf", page_images=20)
        assert result.kind is ExtractionKind.TEXT
        assert "testo vero" in result.page_texts[0]
        assert result.page_texts[1:] == ["", ""]
        assert [image.page for image in result.page_images] == [2, 3]
        assert result.pages == 3
        # Partly unread, whatever the cover sheet said.
        assert not result.inspected

    def test_a_page_with_a_few_characters_is_still_a_picture(self) -> None:
        """A page number on a scan is not a text layer."""
        pdf = mixed_pdf(pdf_bytes("12"), scan_pdf(1, (200, 300)))
        result = extract(pdf, media_type="application/pdf", page_images=20)
        assert [image.page for image in result.page_images] == [1, 2]

    def test_the_page_cap_truncates_and_says_so(self) -> None:
        result = extract(scan_pdf(55, (200, 300)), media_type="application/pdf", page_images=99)
        assert len(result.page_images) == 50  # clamped to the server cap
        assert result.pages == 55
        assert result.images_truncated is True

    def test_a_smaller_request_is_honoured(self) -> None:
        result = extract(scan_pdf(5, (200, 300)), media_type="application/pdf", page_images=2)
        assert [image.page for image in result.page_images] == [1, 2]
        assert result.images_truncated is True

    def test_the_payload_cap_stops_early_and_marks_truncation(
        self, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        import llmp_redaction.documents as documents

        one = extract(scan_pdf(1, (200, 300)), media_type="application/pdf", page_images=1)
        monkeypatch.setattr(documents, "MAX_PAGE_IMAGE_BYTES", len(one.page_images[0].data) * 2 + 1)
        result = extract(scan_pdf(5, (200, 300)), media_type="application/pdf", page_images=5)
        assert len(result.page_images) == 2
        assert result.images_truncated is True

    def test_a_text_pdf_is_read_as_before(self) -> None:
        result = extract(pdf_bytes(LONG_TEXT), media_type="application/pdf", page_images=20)
        assert result.kind is ExtractionKind.TEXT
        assert result.extractor == "markitdown"
        assert result.page_images == [] and result.inspected

    def test_a_scan_that_cannot_be_rendered_says_why(
        self, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        import llmp_redaction.documents as documents

        def fail(*_args: object, **_kwargs: object) -> None:
            raise documents._RenderFailed

        monkeypatch.setattr(documents, "_pdf_by_page", fail)
        result = extract(scan_pdf(1, (200, 300)), media_type="application/pdf", page_images=5)
        assert result.kind is ExtractionKind.NO_TEXT_LAYER
        assert result.page_images == []
        assert "could not be rendered" in result.detail
